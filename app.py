import csv
import io
import threading
from pathlib import Path
from zipfile import ZipFile
from typing import Annotated, Dict, Union
from uuid import uuid4

import fitz
from openpyxl import load_workbook

import qa_review
from qa_review import extract_date_from_filename
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles

BASE_DIR = Path(__file__).parent
STATIC_DIR = BASE_DIR / "static"
MAX_FILE_SIZE = 500 * 1024 * 1024
PREVIEW_ROW_LIMIT = 5000
PREVIEW_COLUMN_LIMIT = 60

app = FastAPI(title="PDF Compare")
files: Dict[str, Dict[str, Union[bytes, str, int]]] = {}

# Opening a PDF from raw bytes reparses the whole file, which is expensive
# for large or malformed documents. Since the frontend requests every page
# of a document at once, doing that reparse per-page multiplies the cost by
# the page count. Caching the opened fitz.Document (and the lock guarding
# it, since a single fitz.Document isn't safe to use from multiple threads
# concurrently) means the expensive parse happens once per document instead
# of once per page. Rendered page PNGs are cached too so re-viewing a page
# (re-opening a panel, switching zoom) is instant.
_pdf_cache: Dict[str, fitz.Document] = {}
_pdf_locks: Dict[str, threading.Lock] = {}
_pdf_cache_guard = threading.Lock()
_page_cache: Dict[str, Dict[tuple, bytes]] = {}


def _get_cached_pdf(document_id: str, content: bytes) -> tuple[fitz.Document, threading.Lock]:
    with _pdf_cache_guard:
        lock = _pdf_locks.setdefault(document_id, threading.Lock())
    with lock:
        pdf = _pdf_cache.get(document_id)
        if pdf is None:
            pdf = fitz.open(stream=content, filetype="pdf")
            _pdf_cache[document_id] = pdf
    return pdf, lock


def _drop_cached_pdf(document_id: str) -> None:
    with _pdf_cache_guard:
        lock = _pdf_locks.pop(document_id, None)
    pdf = _pdf_cache.pop(document_id, None)
    _page_cache.pop(document_id, None)
    if pdf is not None:
        lock_ctx = lock or threading.Lock()
        with lock_ctx:
            pdf.close()

KNOWN_APP_FOLDERS = {
    "artifacts",
    "capital_one_cards",
    "cash_app_p2p",
    "experian",
    "irs_gov",
    "navy_federal_cu",
    "rocket_mortgage",
}


@app.get("/api/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


def register_document(filename: str, content: bytes) -> Dict[str, Union[str, int]]:
    extension = Path(filename).suffix.lower()
    if extension not in {".pdf", ".xlsx", ".csv"}:
        raise ValueError("Only PDF, XLSX, and CSV files are supported.")
    if not content:
        raise ValueError("The file is empty.")
    if len(content) > MAX_FILE_SIZE:
        raise ValueError("Files must be 500 MB or smaller.")

    document_id = uuid4().hex
    kind = "pdf" if extension == ".pdf" else extension[1:]
    page_count = 0
    if kind == "pdf":
        try:
            pdf = fitz.open(stream=content, filetype="pdf")
            page_count = len(pdf)
            pdf.close()
        except Exception as exc:
            raise ValueError(f"Could not read PDF: {exc}") from exc
    else:
        try:
            if kind == "xlsx":
                workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
                workbook.close()
            else:
                content.decode("utf-8-sig")
        except Exception as exc:
            raise ValueError(f"Could not read spreadsheet: {exc}") from exc

    files[document_id] = {
        "name": filename,
        "content": content,
        "kind": kind,
        "page_count": page_count,
    }
    return {
        "id": document_id,
        "name": files[document_id]["name"],
        "size": len(content),
        "kind": kind,
        "page_count": page_count,
    }


@app.post("/api/documents")
async def upload_document(file: Annotated[UploadFile, File(...)]) -> Dict[str, Union[str, int]]:
    filename = file.filename or "Untitled"
    content = await file.read()
    try:
        return register_document(filename, content)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


def determine_folder_name(relative_path: Path) -> str:
    # Prefer a directory segment matching a known app name, wherever
    # it sits in the path, so files are grouped consistently even
    # if a case/wrapper folder sits above the app folder at an
    # inconsistent depth.
    folder_name = next(
        (part for part in relative_path.parts[:-1] if part.lower() in KNOWN_APP_FOLDERS),
        None,
    )
    if folder_name is None:
        # No recognized app name in the path -- fall back to the
        # file's immediate containing folder, whatever it's
        # called, so distinct subfolders always end up as
        # distinct columns instead of collapsing into one.
        folder_name = relative_path.parts[-2] if len(relative_path.parts) >= 2 else "Root files"
    return folder_name


def group_entries_into_folders(entries: list[tuple[str, bytes]]) -> dict:
    """Group (relative_path, content) pairs into app folders the same way
    regardless of whether they came from inside a ZIP or from a folder the
    OS already extracted -- both are just a set of relative paths."""
    folders: Dict[str, list] = {}
    for raw_path, content in entries:
        relative_path = Path(raw_path)
        if relative_path.is_absolute() or ".." in relative_path.parts:
            continue
        if relative_path.suffix.lower() not in {".pdf", ".xlsx", ".csv"}:
            continue
        folder_name = determine_folder_name(relative_path)
        try:
            document = register_document(relative_path.name, content)
        except ValueError:
            continue
        document["folder"] = folder_name
        folders.setdefault(folder_name, []).append(document)

    if not folders:
        raise HTTPException(status_code=422, detail="No supported PDF, CSV, or XLSX files were found.")

    for documents in folders.values():
        dates = {doc["id"]: extract_date_from_filename(str(doc["name"])) for doc in documents}
        dated = [doc for doc in documents if dates[doc["id"]] is not None]
        undated = [doc for doc in documents if dates[doc["id"]] is None]
        dated.sort(key=lambda doc: dates[doc["id"]])
        documents[:] = dated + undated
    return {"folders": [{"name": name, "documents": documents} for name, documents in folders.items()]}


def build_qa_app_entries(all_entries: list[tuple[str, bytes]]) -> Dict[str, list]:
    """Group every non-control file (any extension, not just pdf/csv/xlsx)
    by app folder, for the QA review's file-count and hygiene checks --
    those need to see everything that was actually exported, not just the
    subset this tool knows how to preview."""
    app_entries: Dict[str, list] = {}
    for raw_path, content in all_entries:
        relative_path = Path(raw_path)
        if relative_path.is_absolute() or ".." in relative_path.parts:
            continue
        if relative_path.name in qa_review.CONTROL_FILENAMES:
            continue
        folder_name = determine_folder_name(relative_path)
        app_entries.setdefault(folder_name, []).append((relative_path.name, content))
    return app_entries


@app.post("/api/archive")
async def upload_archive(file: Annotated[UploadFile, File(...)]) -> dict:
    if not (file.filename or "").lower().endswith(".zip"):
        raise HTTPException(status_code=415, detail="Only ZIP files are supported.")
    archive = await file.read()
    if len(archive) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail="ZIP files must be 500 MB or smaller.")

    try:
        with ZipFile(io.BytesIO(archive)) as zip_file:
            all_entries = [
                (entry.filename, zip_file.read(entry))
                for entry in zip_file.infolist()
                if not entry.is_dir()
            ]
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Could not read ZIP: {exc}") from exc

    entries = [
        (raw_path, content) for raw_path, content in all_entries
        if Path(raw_path).suffix.lower() in {".pdf", ".xlsx", ".csv"}
    ]
    result = group_entries_into_folders(entries)
    result["qa_report"] = qa_review.run_phase1_review(all_entries, build_qa_app_entries(all_entries))
    return result


@app.post("/api/folder")
async def upload_folder(files: Annotated[list[UploadFile], File(...)]) -> dict:
    total_size = 0
    entries: list[tuple[str, bytes]] = []
    for file in files:
        content = await file.read()
        total_size += len(content)
        if total_size > MAX_FILE_SIZE:
            raise HTTPException(status_code=413, detail="A folder upload must be 500 MB or smaller in total.")
        entries.append((file.filename or "", content))
    return group_entries_into_folders(entries)


@app.get("/api/documents/{document_id}/preview")
def get_spreadsheet_preview(document_id: str) -> dict:
    document = files.get(document_id)
    if document is None:
        raise HTTPException(status_code=404, detail="Document not found.")
    kind = document["kind"]
    if kind == "pdf":
        raise HTTPException(status_code=400, detail="PDFs do not have a spreadsheet preview.")

    if kind == "csv":
        source = io.StringIO(document["content"].decode("utf-8-sig"), newline="")
        rows = []
        truncated = False
        for row in csv.reader(source):
            if len(rows) >= PREVIEW_ROW_LIMIT:
                truncated = True
                break
            rows.append(row)
        return {"sheets": [{
            "name": "CSV",
            "rows": rows,
            "truncated": truncated,
            "total_rows": None if truncated else len(rows),
        }]}

    try:
        workbook = load_workbook(io.BytesIO(document["content"]), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Could not read spreadsheet: {exc}") from exc

    sheets = []
    try:
        for worksheet in workbook.worksheets:
            rows = []
            truncated = False
            try:
                for row in worksheet.iter_rows(values_only=True):
                    if len(rows) >= PREVIEW_ROW_LIMIT:
                        truncated = True
                        break
                    rows.append(["" if value is None else str(value) for value in row[:PREVIEW_COLUMN_LIMIT]])
            except Exception as exc:
                rows.append([f"Preview stopped: {exc}"])
            sheets.append({"name": worksheet.title, "rows": rows, "truncated": truncated, "total_rows": None if truncated else len(rows)})
    finally:
        workbook.close()
    return {"sheets": sheets}


DOWNLOAD_MEDIA_TYPES = {
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv",
}


@app.get("/api/documents/{document_id}")
def get_document(document_id: str) -> Response:
    document = files.get(document_id)
    if document is None:
        raise HTTPException(status_code=404, detail="File not found.")
    media_type = DOWNLOAD_MEDIA_TYPES.get(document["kind"], "application/octet-stream")
    safe_name = str(document["name"]).replace('"', "'")
    return Response(
        content=document["content"],
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{safe_name}"'},
    )


@app.get("/api/documents/{document_id}/pages/{page_number}")
def get_page(document_id: str, page_number: int, rotation: int = 0) -> Response:
    document = files.get(document_id)
    if document is None:
        raise HTTPException(status_code=404, detail="PDF not found.")
    page_count = int(document["page_count"])
    if page_number < 0 or page_number >= page_count:
        raise HTTPException(status_code=404, detail="Page not found.")
    rotation = rotation % 360
    if rotation not in (0, 90, 180, 270):
        raise HTTPException(status_code=400, detail="Rotation must be a multiple of 90 degrees.")

    cache_key = (page_number, rotation)
    cached_page = _page_cache.get(document_id, {}).get(cache_key)
    if cached_page is not None:
        return Response(content=cached_page, media_type="image/png")

    pdf, lock = _get_cached_pdf(document_id, document["content"])
    with lock:
        page = pdf[page_number]
        page.set_rotation(rotation)
        pixmap = page.get_pixmap(matrix=fitz.Matrix(2.5, 2.5), alpha=False)
        png_bytes = pixmap.tobytes("png")
    _page_cache.setdefault(document_id, {})[cache_key] = png_bytes
    return Response(content=png_bytes, media_type="image/png")


@app.get("/api/documents/{document_id}/text")
def get_pdf_text(document_id: str) -> dict[str, str]:
    document = files.get(document_id)
    if document is None:
        raise HTTPException(status_code=404, detail="PDF not found.")
    if document["kind"] != "pdf":
        raise HTTPException(status_code=400, detail="Only PDFs support text extraction.")

    pdf, lock = _get_cached_pdf(document_id, document["content"])
    with lock:
        text = "\n\n".join(page.get_text().strip() for page in pdf)
    return {"text": text}


@app.delete("/api/documents/{document_id}")
def delete_document(document_id: str) -> dict[str, bool]:
    _drop_cached_pdf(document_id)
    if files.pop(document_id, None) is None:
        raise HTTPException(status_code=404, detail="PDF not found.")
    return {"deleted": True}


@app.get("/")
def index() -> FileResponse:
    return FileResponse(STATIC_DIR / "index.html", headers={"Cache-Control": "no-cache"})


app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


@app.middleware("http")
async def no_cache_static(request, call_next):
    response = await call_next(request)
    if request.url.path.startswith("/static/"):
        response.headers["Cache-Control"] = "no-cache"
    return response
